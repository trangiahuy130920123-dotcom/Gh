import io
import json
import os
import streamlit as st
from docx import Document
from pptx import Presentation
from pptx.util import Inches, Pt
from google import genai

# ---------------------------------------------------------
# 1. CẤU HÌNH TRANG STREAMLIT
# ---------------------------------------------------------
st.set_page_config(
    page_title="Nova AI - Trợ lý & Tạo Slide Thông Minh",
    page_icon="✨",
    layout="centered",
    initial_sidebar_state="expanded"
)

# ---------------------------------------------------------
# 2. CSS TÙY CHỈNH GIAO DIỆN
# ---------------------------------------------------------
st.markdown("""
<style>
    header { visibility: hidden; }
    footer { visibility: hidden; }
    #MainMenu { visibility: hidden; }
    
    .block-container {
        padding-top: 1.5rem !important;
        padding-bottom: 4rem !important;
        max-width: 720px;
    }

    .hero-title {
        text-align: center;
        font-size: 2.2rem;
        font-weight: 800;
        color: #0f172a;
        margin-top: 10px;
        margin-bottom: 8px;
        line-height: 1.2;
    }
    
    .hero-sub {
        text-align: center;
        color: #64748b;
        font-size: 0.95rem;
        margin-bottom: 1.8rem;
        line-height: 1.5;
    }

    div.stButton > button {
        width: 100%;
        border-radius: 12px;
        border: 1px solid #e2e8f0;
        padding: 12px 16px;
        background-color: #ffffff;
        box-shadow: 0 2px 4px rgba(0,0,0,0.02);
        transition: all 0.2s ease;
    }
    
    div.stButton > button:hover {
        border-color: #6366f1;
        background-color: #f8fafc;
    }
</style>
""", unsafe_allow_html=True)

# ---------------------------------------------------------
# 3. PHẦN CÀI ĐẶT API KEY (SIDEBAR)
# ---------------------------------------------------------
with st.sidebar:
    st.title("⚙️ Cấu hình AI")
    
    env_api_key = os.environ.get("GEMINI_API_KEY", "")
    if "GEMINI_API_KEY" in st.secrets:
        env_api_key = st.secrets["GEMINI_API_KEY"]

    api_key_input = st.text_input(
        "Nhập Google Gemini API Key:",
        value=env_api_key,
        type="password",
        help="Lấy API Key miễn phí tại https://aistudio.google.com/"
    )
    
    st.markdown("""
    ---
    💡 **Hướng dẫn lấy API Key miễn phí:**
    1. Truy cập [Google AI Studio](https://aistudio.google.com/)
    2. Đăng nhập Google & bấm **Get API key**
    3. Dán mã Key vào ô trên để kích hoạt AI thật!
    """)

# Khởi tạo Gemini Client
client = None
if api_key_input.strip():
    try:
        client = genai.Client(api_key=api_key_input.strip())
    except Exception as e:
        st.sidebar.error(f"Lỗi khởi tạo API: {e}")

# ---------------------------------------------------------
# 4. HÀM TẠO FILE XUẤT (PPTX, DOCX, PYTHON)
# ---------------------------------------------------------
def generate_pptx(topic, slides_data):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]
    title_slide = prs.slides.add_slide(blank_layout)
    tx_box = title_slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11.333), Inches(2))
    tf = tx_box.text_frame
    p = tf.paragraphs[0]
    p.text = topic
    p.font.bold = True
    p.font.size = Pt(44)
    p.font.name = "Arial"

    for slide_item in slides_data:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        title_shape = slide.shapes.title
        title_shape.text = slide_item.get("title", "Slide")
        
        body_shape = slide.placeholders[1]
        tf_body = body_shape.text_frame
        tf_body.word_wrap = True
        
        bullets = slide_item.get("bullets", [])
        for idx, bullet in enumerate(bullets):
            if idx == 0:
                p_item = tf_body.paragraphs[0]
                p_item.text = bullet
            else:
                p_item = tf_body.add_paragraph()
                p_item.text = bullet

    buffer = io.BytesIO()
    prs.save(buffer)
    buffer.seek(0)
    return buffer


def generate_docx(topic, slides_data):
    doc = Document()
    doc.add_heading(topic, 0)

    for slide in slides_data:
        doc.add_heading(slide.get("title", "Slide"), level=1)
        for bullet in slide.get("bullets", []):
            doc.add_paragraph(bullet, style='List Bullet')

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def generate_python_script(topic, slides_data):
    slides_json = json.dumps(slides_data, ensure_ascii=False, indent=4)
    topic_json = json.dumps(topic, ensure_ascii=False)
    
    script_content = f"""# -*- coding: utf-8 -*-
\"\"\"
Nova AI Generated Presentation Script
Topic: {topic}
\"\"\"

topic = {topic_json}
slides = {slides_json}

def main():
    print(f"=== BÀI THUYẾT TRÌNH: {{topic}} ===")
    for idx, slide in enumerate(slides, 1):
        print(f"\\n[Slide {{idx}}] {{slide.get('title', '')}}")
        for item in slide.get("bullets", []):
            print(f"  - {{item}}")

if __name__ == "__main__":
    main()
"""
    return script_content.encode("utf-8")

# ---------------------------------------------------------
# 5. INITIALIZE SESSION STATE
# ---------------------------------------------------------
if "mode" not in st.session_state:
    st.session_state.mode = "presentation"

if "messages" not in st.session_state:
    st.session_state.messages = []

if "generated_slides" not in st.session_state:
    st.session_state.generated_slides = None

if "current_topic" not in st.session_state:
    st.session_state.current_topic = ""

# ---------------------------------------------------------
# 6. HEADER & ĐIỀU HƯỚNG
# ---------------------------------------------------------
st.markdown('<div class="hero-title">What can I help you create?</div>', unsafe_allow_html=True)
st.markdown('<div class="hero-sub">Hỏi AI bất cứ điều gì — hoặc yêu cầu tạo bài trình bày và xuất file PowerPoint, Word, Python.</div>', unsafe_allow_html=True)

col1, col2 = st.columns(2)

with col1:
    if st.button("💬 **Hỏi đáp AI**\n\nHỏi đáp trực tiếp cùng mô hình Gemini AI."):
        st.session_state.mode = "chat"

with col2:
    if st.button("📊 **Tạo trình bày**\n\nTạo slide thuyết trình AI & Xuất tệp PPTX/Word/Python."):
        st.session_state.mode = "presentation"

st.divider()

# ---------------------------------------------------------
# 7. CHẾ ĐỘ 1: TẠO SLIDE
# ---------------------------------------------------------
if st.session_state.mode == "presentation":
    st.subheader("📊 Tạo bài thuyết trình thông minh bằng AI")
    
    topic_input = st.text_input(
        "Chủ đề bài thuyết trình:", 
        placeholder="Ví dụ: Ứng dụng AI trong giáo dục..."
    )
    
    num_slides = st.slider("Số lượng Slide:", min_value=3, max_value=10, value=4)

    if st.button("🚀 Bắt đầu tạo Slide bằng AI", type="primary"):
        if not topic_input.strip():
            st.warning("Vui lòng nhập chủ đề bài thuyết trình.")
        else:
            st.session_state.current_topic = topic_input.strip()
            
            if client:
                with st.spinner("Gemini AI đang tư duy và lập dàn ý slide..."):
                    try:
                        prompt = f"""Bạn là một chuyên gia thuyết trình. Hãy tạo bài thuyết trình gồm {num_slides} slide cho chủ đề: "{topic_input}".
Yêu cầu trả về đúng định dạng JSON dạng danh sách các object:
[
  {{"title": "1. Giới thiệu", "bullets": ["Ý 1", "Ý 2", "Ý 3"]}},
  {{"title": "2. Thách thức", "bullets": ["Ý 1", "Ý 2", "Ý 3"]}}
]
Chỉ trả về chuỗi JSON thuần túy, không kèm mã codeblock hay văn bản khác.
"""
                        response = client.models.generate_content(
                            model="gemini-2.5-flash",
                            contents=prompt
                        )
                        
                        raw_text = response.text.strip()
                        if raw_text.startswith("```"):
                            raw_text = raw_text.split("```")[1]
                            if raw_text.startswith("json"):
                                raw_text = raw_text[4:]
                        raw_text = raw_text.strip()
                        
                        slides_json = json.loads(raw_text)
                        st.session_state.generated_slides = slides_json
                        st.success("🎉 Gemini AI đã khởi tạo thành công bài thuyết trình!")
                    except Exception as e:
                        st.error(f"Lỗi khi gọi Gemini AI: {e}")
            else:
                st.info("💡 Chưa nhập Gemini API Key. Hệ thống đang tạo bản mẫu thử nghiệm.")
                st.session_state.generated_slides = [
                    {"title": f"1. Giới thiệu về {topic_input}", "bullets": ["Khái niệm tổng quan.", "Xu hướng phát triển hiện tại.", "Tầm quan trọng của chủ đề."]},
                    {"title": "2. Các giá trị cốt lõi", "bullets": ["Tăng hiệu suất công việc.", "Tối ưu hóa quy trình.", "Đổi mới sáng tạo."]},
                    {"title": "3. Kết luận & Q&A", "bullets": ["Tóm tắt nội dung chính.", "Đề xuất hành động tiếp theo.", "Giải đáp thắc mắc."]}
                ][:num_slides]

    if st.session_state.generated_slides:
        st.write("---")
        st.markdown(f"### 📋 Xem trước nội dung: **{st.session_state.current_topic}**")

        for idx, slide in enumerate(st.session_state.generated_slides, 1):
            with st.expander(f"Slide {idx}: {slide.get('title', '')}", expanded=True):
                for bullet in slide.get("bullets", []):
                    st.write(f"• {bullet}")

        st.markdown("### 📥 Xuất Tệp Trình Bày")
        exp_col1, exp_col2, exp_col3 = st.columns(3)

        with exp_col1:
            pptx_file = generate_pptx(st.session_state.current_topic, st.session_state.generated_slides)
            st.download_button(
                label="📄 PowerPoint (.pptx)",
                data=pptx_file,
                file_name=f"{st.session_state.current_topic}.pptx",
                mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                use_container_width=True
            )

        with exp_col2:
            docx_file = generate_docx(st.session_state.current_topic, st.session_state.generated_slides)
            st.download_button(
                label="📝 Word (.docx)",
                data=docx_file,
                file_name=f"{st.session_state.current_topic}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                use_container_width=True
            )

        with exp_col3:
            py_file = generate_python_script(st.session_state.current_topic, st.session_state.generated_slides)
            st.download_button(
                label="🐍 Python Script (.py)",
                data=py_file,
                file_name=f"{st.session_state.current_topic}.py",
                mime="text/x-python",
                use_container_width=True
            )

# ---------------------------------------------------------
# 8. CHẾ ĐỘ 2: HỎI ĐÁP CHAT
# ---------------------------------------------------------
elif st.session_state.mode == "chat":
    st.subheader("💬 Trò chuyện trực tiếp cùng Gemini AI")

    for message in st.session_state.messages:
        with st.chat_message(message["role"]):
            st.markdown(message["content"])

    if prompt := st.chat_input("Hỏi AI bất cứ điều gì..."):
        st.session_state.messages.append({"role": "user", "content": prompt})
        with st.chat_message("user"):
            st.markdown(prompt)

        with st.chat_message("assistant"):
            if client:
                with st.spinner("Gemini AI đang trả lời..."):
                    try:
                        response = client.models.generate_content(
                            model="gemini-2.5-flash",
                            contents=prompt
                        )
                        reply_text = response.text
                    except Exception as e:
                        reply_text = f"Lỗi gọi AI: {e}"
            else:
                reply_text = "🤖 [Chế độ dùng thử] Bạn chưa nhập Gemini API Key ở menu góc trái. Hãy nhập API Key để trò chuyện trực tiếp cùng Gemini AI nhé!"

            st.markdown(reply_text)
            st.session_state.messages.append({"role": "assistant", "content": reply_text})
